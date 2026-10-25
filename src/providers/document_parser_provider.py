import io
import logging
import re
import unicodedata
from pathlib import Path
from typing import Any

import pymupdf

logger = logging.getLogger(__name__)


class DocumentParserProvider:
    SUPPORTED = {
        ".pdf", ".txt", ".docx",
        ".png", ".jpg", ".jpeg", ".tif", ".tiff",
    }

    # Short or sparse PDF text is a signal to check the page with OCR.
    OCR_MIN_CHARACTERS = 180
    OCR_MIN_WORDS = 20
    OCR_RENDER_SCALE = 3

    @staticmethod
    def _word_count(text: str) -> int:
        return len(re.findall(r"\b[\w'-]+\b", text, flags=re.UNICODE))

    @classmethod
    def _needs_ocr(cls, text: str) -> bool:
        cleaned = text.strip()

        if not cleaned:
            return True

        bad = sum(
            1
            for char in cleaned
            if unicodedata.category(char) == "Cc"
            and char not in "\n\r\t"
        )

        if bad > max(3, int(len(cleaned) * 0.02)):
            return True

        return (
            len(cleaned) < cls.OCR_MIN_CHARACTERS
            or cls._word_count(cleaned) < cls.OCR_MIN_WORDS
        )

    @staticmethod
    def _is_ocr_better(native_text: str, ocr_text: str) -> bool:
        native = native_text.strip()
        ocr = ocr_text.strip()

        if not ocr:
            return False

        if not native:
            return True

        native_words = DocumentParserProvider._word_count(native)
        ocr_words = DocumentParserProvider._word_count(ocr)

        native_chars = sum(char.isalnum() for char in native)
        ocr_chars = sum(char.isalnum() for char in ocr)

        # Avoid replacing good native text with a similarly sized OCR result.
        enough_extra_text = (
            len(ocr) >= len(native) * 1.15
            and ocr_words >= native_words + 2
            and ocr_chars > native_chars
        )

        # When native extraction is clearly sparse, prefer a substantially
        # more informative OCR result.
        recovers_missing_content = (
            DocumentParserProvider._needs_ocr(native)
            and ocr_words >= max(native_words + 3, 8)
            and ocr_chars >= max(native_chars + 20, 40)
        )

        return enough_extra_text or recovers_missing_content

    def _ocr_image(self, image: Any) -> str:
        try:
            from rapidocr import RapidOCR
            from rapidocr.utils.typings import (
                EngineType,
                ModelType,
                OCRVersion,
            )

            if not hasattr(self, "_ocr"):
                self._ocr = RapidOCR(params={
                    "Det.engine_type": EngineType.ONNXRUNTIME,
                    "Det.lang_type": "multi",
                    "Det.model_type": ModelType.MOBILE,
                    "Det.ocr_version": OCRVersion.PPOCRV5,
                    "Rec.engine_type": EngineType.ONNXRUNTIME,
                    "Rec.lang_type": "en",
                    "Rec.model_type": ModelType.MOBILE,
                    "Rec.ocr_version": OCRVersion.PPOCRV5,
                })

            result = self._ocr(image)
            lines = getattr(result, "txts", None) or []

            return "\n".join(
                str(line).strip()
                for line in lines
                if str(line).strip()
            )

        except Exception:
            logger.exception(
                "OCR failed. Check RapidOCR configuration and model availability."
            )
            raise

    def extract_pages(self, file_path: str | Path) -> list[dict[str, Any]]:
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix not in self.SUPPORTED:
            raise ValueError(f"Unsupported document extension: {suffix}")

        if suffix == ".txt":
            text = path.read_text(encoding="utf-8-sig", errors="replace")
            return [{
                "page_number": 1,
                "text": text.strip(),
                "extraction_method": "text",
            }]

        if suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(
                        cell.text.strip() for cell in row.cells
                    ))

            return [{
                "page_number": 1,
                "text": "\n".join(parts).strip(),
                "extraction_method": "docx",
            }]

        if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
            from PIL import Image

            with Image.open(path) as image:
                text = self._ocr_image(image.convert("RGB"))

            return [{
                "page_number": 1,
                "text": text,
                "extraction_method": "ocr",
            }]

        pages: list[dict[str, Any]] = []

        with pymupdf.open(path) as pdf:
            page_count = len(pdf)

            for number, page in enumerate(pdf, start=1):
                native_text = page.get_text("text").strip()
                text = native_text
                method = "text"

                if self._needs_ocr(native_text):
                    logger.info(
                        "Page %s has sparse or suspicious text; trying OCR.",
                        number,
                    )

                    pixmap = page.get_pixmap(
                        matrix=pymupdf.Matrix(
                            self.OCR_RENDER_SCALE,
                            self.OCR_RENDER_SCALE,
                        ),
                        alpha=False,
                    )

                    from PIL import Image

                    image = Image.open(
                        io.BytesIO(pixmap.tobytes("png"))
                    ).convert("RGB")

                    ocr_text = self._ocr_image(image).strip()

                    if self._is_ocr_better(native_text, ocr_text):
                        text = ocr_text
                        method = "ocr"
                        logger.info(
                            "Page %s: using OCR text (%s chars vs %s).",
                            number,
                            len(ocr_text),
                            len(native_text),
                        )
                    else:
                        logger.info(
                            "Page %s: keeping native text; OCR was not better.",
                            number,
                        )

                pages.append({
                    "page_number": number,
                    "text": text,
                    "extraction_method": method,
                })

        if len(pages) != page_count:
            raise RuntimeError(
                "PDF page extraction did not return all pages."
            )

        return pages
